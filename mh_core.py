"""Core logic for the Sian3L Mode Hopping tool.

1. load_inputs()      : read csv / xlsx / zip (any folder depth) -> raw tables per SN
2. build_curve_xlsx() : paste raw data into Sian3L_MH_Chart-Template.xlsx (XML level, charts kept)
3. build_ppt()        : generate the PPT report from Sian3L_Mode_Hopping_Template.pptx
"""
from __future__ import annotations

import copy
import io
import os
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date
from xml.sax.saxutils import escape

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import openpyxl  # noqa: E402
import pandas as pd  # noqa: E402
from openpyxl.utils import get_column_letter as col_letter  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Pt  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
def _tpl(name):
    """Template in templates/ or next to this file (repo root)."""
    for p in (os.path.join(HERE, "templates", name), os.path.join(HERE, name)):
        if os.path.exists(p):
            return p
    return os.path.join(HERE, name)


XLSX_TEMPLATE = _tpl("Sian3L_MH_Chart-Template.xlsx")
PPT_TEMPLATE = _tpl("Sian3L_Mode_Hopping_Template.pptx")

HEADERS = ["[CHNumber]", "[Times]", "[Case_Temp]", "[Wavelength]", "[TXP]",
           "[CH1_DDMI_Bias]", "[CH2_DDMI_Bias]", "[CH3_DDMI_Bias]", "[CH4_DDMI_Bias]",
           "[CH5_DDMI_Bias]", "[CH6_DDMI_Bias]", "[CH7_DDMI_Bias]", "[CH8_DDMI_Bias]",
           "[dW/dT]", "[Peak_WL_Jump_Max]", "[Peak_WL_Jump_Min]", "[Peak_WL_Jump_Avg]",
           "[CH_Pass_Fail]", "[FailureCodeID]", "[ErrorMessage]"]
NCOL = len(HEADERS)                     # A..T
DWDT_MIN, DWDT_MAX = 0.0, 0.2           # spec lines (Curve!T / Curve!S)

# template sheet files (see xl/workbook.xml of the template)
SHEET_CURVE = "xl/worksheets/sheet1.xml"
SHEET_SPLIT = {"xl/worksheets/sheet2.xml": "Operational", "xl/worksheets/sheet3.xml": "Operational",
               "xl/worksheets/sheet5.xml": "Maximum", "xl/worksheets/sheet6.xml": "Maximum"}
SHEET_RAW = {"Operational": "xl/worksheets/sheet4.xml", "Maximum": "xl/worksheets/sheet7.xml"}


# --------------------------------------------------------------------------- input
@dataclass
class SNData:
    sn: str
    date: str = ""                                   # MM-DD from file name
    files: dict = field(default_factory=dict)        # kind -> file name
    raw: dict = field(default_factory=dict)          # kind -> list[list] (20 cols, no header)

    def channel(self, kind: str, ch: str) -> pd.DataFrame:
        rows = [r for r in self.raw.get(kind, []) if str(r[0]) == f"{ch}_{kind}"]
        return pd.DataFrame(rows, columns=HEADERS)


NAME_RE = re.compile(r"^(?P<sn>.+?)_(?P<kind>Operational|Maximum)[ _]?Current(?:_(?P<date>\d{2}-\d{2}))?",
                     re.IGNORECASE)


def _clean(v):
    if v is None:
        return None
    if isinstance(v, float) and pd.isna(v):
        return None
    if isinstance(v, str):
        s = v.strip()
        if s == "":
            return None
        try:
            f = float(s)
            return int(f) if f.is_integer() and "." not in s else f
        except ValueError:
            return s
    if hasattr(v, "item"):          # numpy scalar
        return v.item()
    return v


def _rows_from_table(table: list[list]) -> list[list] | None:
    """Find the [CHNumber] header row and return the data rows (20 cols)."""
    for i, row in enumerate(table[:20]):
        cells = [str(c).strip() if c is not None else "" for c in row]
        if cells and cells[0] == "[CHNumber]":
            idx = {h: cells.index(h) for h in HEADERS if h in cells}
            out = []
            for r in table[i + 1:]:
                vals = [_clean(r[idx[h]]) if h in idx and idx[h] < len(r) else None for h in HEADERS]
                if vals[0] is None:
                    continue
                out.append(vals)
            return out
    return None


def _read_table(name: str, data: bytes) -> list[list] | None:
    ext = os.path.splitext(name)[1].lower()
    if ext == ".csv":
        for enc in ("utf-8-sig", "big5", "cp950", "latin-1"):
            try:
                df = pd.read_csv(io.BytesIO(data), header=None, dtype=str, encoding=enc,
                                 keep_default_na=False, on_bad_lines="skip")
                return _rows_from_table(df.values.tolist())
            except UnicodeDecodeError:
                continue
        return None
    if ext in (".xlsx", ".xlsm"):
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for ws in wb.worksheets:                     # first sheet that has a [CHNumber] header
            rows = _rows_from_table([list(r) for r in ws.iter_rows(values_only=True)])
            if rows:
                return rows
    return None


def _walk(name: str, data: bytes):
    """Yield (file name, bytes) for every csv/xlsx, opening zips (and zips in zips) at any depth."""
    base = os.path.basename(name)
    if base.startswith(("~$", "._")) or "__MACOSX" in name:
        return
    ext = os.path.splitext(name)[1].lower()
    if ext == ".zip":
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            for info in z.infolist():
                if not info.is_dir():
                    yield from _walk(info.filename, z.read(info))
    elif ext in (".csv", ".xlsx", ".xlsm"):
        yield base, data


def load_inputs(files: list[tuple[str, bytes]]) -> tuple[dict[str, SNData], list[str]]:
    """files: [(name, bytes)] -> ({sn: SNData}, warnings)"""
    sns: dict[str, SNData] = {}
    warnings = []
    for name, data in files:
        for fname, fdata in _walk(name, data):
            m = NAME_RE.match(os.path.splitext(fname)[0])
            try:
                rows = _read_table(fname, fdata)
            except Exception as e:  # noqa: BLE001
                warnings.append(f"{fname}: cannot read ({e})")
                continue
            if not rows:
                warnings.append(f"{fname}: no [CHNumber] raw-data table found - skipped")
                continue
            kinds = {str(r[0]).split("_", 1)[-1].capitalize() for r in rows}
            kind = (m.group("kind").capitalize() if m else None) or (kinds.pop() if len(kinds) == 1 else None)
            if kind not in ("Operational", "Maximum"):
                warnings.append(f"{fname}: cannot tell Operational / Maximum - skipped")
                continue
            sn = m.group("sn") if m else os.path.splitext(fname)[0]
            d = sns.setdefault(sn, SNData(sn))
            if kind in d.raw:
                warnings.append(f"{sn}: more than one {kind} file - using {fname}")
            d.raw[kind] = rows
            d.files[kind] = fname
            if m and m.group("date"):
                d.date = m.group("date")
    for sn, d in sns.items():
        for k in ("Operational", "Maximum"):
            if k not in d.raw:
                warnings.append(f"{sn}: missing {k} Current file")
    return dict(sorted(sns.items(), key=lambda kv: _natkey(kv[0]))), warnings


def _natkey(s):
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", s)]


# --------------------------------------------------------------------------- analysis
def analyze(d: SNData) -> dict:
    """Simple per-SN judgement used for the PPT conclusions."""
    res = {"dwdt": {}, "fails": [], "dwdt_pass": True}
    for kind in ("Operational", "Maximum"):
        for ch in ("1", "5"):
            df = d.channel(kind, ch)
            if df.empty:
                continue
            dw = pd.to_numeric(df["[dW/dT]"], errors="coerce").dropna()
            if len(dw):
                lo, hi = dw.min(), dw.max()
                res["dwdt"][(kind, ch)] = (lo, hi)
                if lo < DWDT_MIN or hi > DWDT_MAX:
                    res["dwdt_pass"] = False
            bad = df[df["[CH_Pass_Fail]"].astype(str).str.lower() == "fail"]
            if len(bad):
                t = pd.to_numeric(bad["[Case_Temp]"], errors="coerce")
                msgs = sorted({re.sub(r"\s*\(.*\)", "", str(m)) for m in bad["[ErrorMessage]"] if m not in (None, 0, "0")})
                at = f"{t.min():.1f} ℃" if t.min() == t.max() else f"{t.min():.1f}~{t.max():.1f} ℃"
                res["fails"].append(f"CH{ch} {kind}: {', '.join(msgs) or 'Fail'} at {at} ({len(bad)} pts)")
    return res


def default_conclusions(d: SNData) -> tuple[str, str]:
    a = analyze(d)
    rng = [f"{v[0]:.2f}~{v[1]:.2f}" for v in a["dwdt"].values()]
    span = (f"{min(v[0] for v in a['dwdt'].values()):.2f}~{max(v[1] for v in a['dwdt'].values()):.2f}"
            if rng else "n/a")
    dw = (f"dWL/dT {span} within spec ({DWDT_MIN:g}~{DWDT_MAX:g}), no mode hopping" if a["dwdt_pass"]
          else f"dWL/dT {span} out of spec ({DWDT_MIN:g}~{DWDT_MAX:g}), suspect mode hopping")
    wl = "; ".join(a["fails"]) if a["fails"] else "Wavelength and TxP PASS"
    return wl, dw


# --------------------------------------------------------------------------- 1. Curve xlsx
def _cell(ref, v):
    if v is None or v == "":
        return ""
    if isinstance(v, bool):
        return f'<c r="{ref}" t="b"><v>{int(v)}</v></c>'
    if isinstance(v, (int, float)):
        return f'<c r="{ref}"><v>{repr(v)}</v></c>'
    return f'<c r="{ref}" t="inlineStr"><is><t>{escape(str(v))}</t></is></c>'


def _put_raw(xml, rows):
    head = re.search(r'<row r="1".*?</row>', xml, re.S).group(0)
    body = "".join(
        f'<row r="{i}" spans="1:{NCOL}">' + "".join(_cell(f"{col_letter(j + 1)}{i}", v) for j, v in enumerate(r)) + "</row>"
        for i, r in enumerate(rows, start=2))
    xml = re.sub(r"<sheetData>.*</sheetData>", lambda m: f"<sheetData>{head}{body}</sheetData>", xml, flags=re.S)
    last = len(rows) + 1
    xml = re.sub(r'<dimension ref="[^"]*"/>', f'<dimension ref="A1:T{last}"/>', xml)
    return re.sub(r'(<autoFilter ref=")[^"]*"', rf'\g<1>A1:T{last}"', xml)


def _strip_cache(xml):
    """Drop cached formula results so Excel shows fresh values; blank rows -> #N/A (charts skip them)."""
    xml = xml.replace('),"",INDEX(', '),NA(),INDEX(')

    def fix(m):
        attrs = re.sub(r'\s+t="[^"]*"', "", m.group(1))
        return f"<c{attrs}>{m.group(2)}</c>"
    return re.sub(r'<c((?:\s+[a-z]+="[^"]*")*)>(<f>.*?</f>)(?:<v>.*?</v>)?</c>', fix, xml, flags=re.S)


def _extend_rows(xml, last_row):
    """Clone the sheet's last row down to last_row (relative refs shifted) if the data is longer."""
    rows = re.findall(r'<row r="(\d+)"', xml)
    cur = max(map(int, rows))
    if last_row <= cur:
        return xml
    tmpl = re.search(rf'<row r="{cur}"[^>]*>.*?</row>', xml, re.S).group(0)
    new = "".join(re.sub(rf"(?<=[A-Z$:]){cur}(?!\d)", str(r), tmpl) for r in range(cur + 1, last_row + 1))
    xml = xml.replace(tmpl, tmpl + new)
    return re.sub(r'(<dimension ref="A1:[A-Z]+)\d+"', rf'\g<1>{last_row}"', xml)


def build_curve_xlsx(d: SNData, template=XLSX_TEMPLATE) -> bytes:
    opr, mx = d.raw.get("Operational", []), d.raw.get("Maximum", [])
    need = {"Operational": max(len(d.channel("Operational", c)) for c in "15") + 1,
            "Maximum": max(len(d.channel("Maximum", c)) for c in "15") + 1}
    curve_last = max(need.values())
    out = io.BytesIO()
    with zipfile.ZipFile(template) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        # template capacity
        cap = {f: max(map(int, re.findall(r'<row r="(\d+)"', zin.read(f).decode()))) for f in SHEET_SPLIT}
        curve_cap = max(map(int, re.findall(r'<row r="(\d+)"', zin.read(SHEET_CURVE).decode())))
        grow = curve_last > curve_cap or any(need[k] > cap[f] for f, k in SHEET_SPLIT.items())
        for it in zin.infolist():
            fn, data = it.filename, zin.read(it.filename)
            if fn == "xl/calcChain.xml":
                continue
            if fn in ("[Content_Types].xml", "xl/_rels/workbook.xml.rels"):
                t = data.decode()
                t = re.sub(r'<Override PartName="/xl/calcChain.xml"[^>]*/>', "", t)
                t = re.sub(r'<Relationship Id="[^"]*" Type="[^"]*calcChain" Target="calcChain.xml"/>', "", t)
                data = t.encode()
            elif fn == SHEET_RAW["Operational"]:
                data = _put_raw(data.decode(), opr).encode()
            elif fn == SHEET_RAW["Maximum"]:
                data = _put_raw(data.decode(), mx).encode()
            elif fn in SHEET_SPLIT:
                t = _strip_cache(data.decode())
                if grow:
                    t = _extend_rows(t, max(need[SHEET_SPLIT[fn]], cap[fn]))
                data = t.encode()
            elif fn == SHEET_CURVE:
                t = _strip_cache(data.decode())
                if grow:
                    t = _extend_rows(t, max(curve_last, curve_cap))
                data = t.encode()
            elif fn.startswith("xl/charts/chart") and fn.endswith(".xml") and grow:
                t = data.decode()
                t = re.sub(r"(Curve!\$[A-Z]+\$2:\$[A-Z]+\$)\d+", rf"\g<1>{max(curve_last, curve_cap)}", t)
                t = re.sub(r"<c:(?:numCache|strCache)>.*?</c:(?:numCache|strCache)>", "", t, flags=re.S)
                data = t.encode()
            elif fn == "xl/workbook.xml":
                t = data.decode()
                t = re.sub(r"<calcPr[^>]*/>", '<calcPr calcId="191029" fullCalcOnLoad="1"/>', t)
                t = re.sub(r"Maximum_RawData!\$A\$1:\$T\$\d+", f"Maximum_RawData!$A$1:$T${len(mx) + 1}", t)
                t = re.sub(r"Operational_RawData!\$A\$1:\$T\$\d+", f"Operational_RawData!$A$1:$T${len(opr) + 1}", t)
                data = t.encode()
            zout.writestr(zipfile.ZipInfo(fn, it.date_time), data, zipfile.ZIP_DEFLATED)
    return out.getvalue()


def curve_filename(d: SNData) -> str:
    return f"{d.sn}_Sian3L_MH_Chart_{d.date or date.today().strftime('%m-%d')}.xlsx"


# --------------------------------------------------------------------------- 2. charts for PPT
C = {"blue": "#4F81BD", "red": "#C0504D", "green": "#9BBB59", "purple": "#8064A2", "aqua": "#4BACC6",
     "gray": "#7F7F7F", "navy": "#1F497D", "spec": "#FF0000"}


def _num(df, col):
    return pd.to_numeric(df[col], errors="coerce")


def _axes(fig_w, fig_h, title, ylab, y2lab=None):
    fig, ax = plt.subplots(figsize=(fig_w, fig_h), dpi=200)
    fig.patch.set_facecolor("white")
    ax.set_title(title, fontsize=11, pad=8)
    ax.set_xlabel("Case_Temp (℃)", fontsize=8)
    ax.set_ylabel(ylab, fontsize=8)
    ax.set_xlim(0, 70)
    ax.set_xticks(range(0, 71, 5))
    ax.grid(axis="y", color="#D9D9D9", linewidth=0.6)
    ax.set_axisbelow(True)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    for s in ("left", "bottom"):
        ax.spines[s].set_color("#BFBFBF")
    ax.tick_params(labelsize=7, colors="#404040", length=0)
    ax2 = None
    if y2lab:
        ax2 = ax.twinx()
        ax2.set_ylabel(y2lab, fontsize=8)
        ax2.tick_params(labelsize=7, colors="#404040", length=0)
        for s in ax2.spines.values():
            s.set_visible(False)
    return fig, ax, ax2


def _finish(fig, ax, ax2=None) -> bytes:
    h, l = ax.get_legend_handles_labels()
    if ax2 is not None:
        h2, l2 = ax2.get_legend_handles_labels()
        h, l = h + h2, l + l2
    fig.legend(h, l, loc="lower center", ncol=len(l), fontsize=7, frameon=False)
    fig.tight_layout(rect=(0, 0.07, 1, 1))
    buf = io.BytesIO()
    fig.savefig(buf, format="png", facecolor="white")
    plt.close(fig)
    return buf.getvalue()


def make_charts(d: SNData, aspect=(2.04, 2.0)) -> dict[str, bytes]:
    """Return the 8 chart PNGs (same series as the Excel template's Curve charts)."""
    o1, o5 = d.channel("Operational", "1"), d.channel("Operational", "5")
    m1, m5 = d.channel("Maximum", "1"), d.channel("Maximum", "5")
    T, WL, TX, dW = "[Case_Temp]", "[Wavelength]", "[TXP]", "[dW/dT]"
    B1, B5 = "[CH1_DDMI_Bias]", "[CH5_DDMI_Bias]"
    W = 6.4
    out = {}
    lw = 1.1

    for ch, o, m, b in (("1", o1, m1, B1), ("5", o5, m5, B5)):
        fig, ax, ax2 = _axes(W, W / aspect[0], f"Wavelength Channel {ch}", "Wavelength (nm)", "Bias(mA)")
        ax.plot(_num(o, T), _num(o, WL), color=C["blue"], lw=lw, label=f"{ch}_Opr_Wavelength")
        ax.plot(_num(m, T), _num(m, WL), color=C["red"], lw=lw, label=f"{ch}_Max_Wavelength")
        ax2.plot(_num(o, T), _num(o, b), color=C["gray"], lw=lw, label=f"[CH{ch}_DDMI_Opr_Bias]")
        ax2.set_ylim(bottom=0)
        out[f"wl{ch}"] = _finish(fig, ax, ax2)

    fig, ax, ax2 = _axes(W, W / aspect[0], "TXP_Operational Bias", "dBm", "Bias(mA)")
    ax.plot(_num(o1, T), _num(o1, TX), color=C["purple"], lw=lw, label="[CH1_Opr_TxP]")
    ax.plot(_num(o5, T), _num(o5, TX), color=C["aqua"], lw=lw, label="[CH5_Opr_TxP]")
    ax2.plot(_num(o1, T), _num(o1, B1), color=C["red"], lw=lw, label="[CH1_DDMI_Opr_Bias]")
    ax2.plot(_num(o5, T), _num(o5, B5), color=C["green"], lw=lw, label="[CH5_DDMI_Opr_Bias]")
    ax.set_ylim(-8, 8)
    ax.set_yticks(range(-8, 9, 2))
    ax2.set_ylim(bottom=80)
    out["txp_opr"] = _finish(fig, ax, ax2)

    fig, ax, ax2 = _axes(W, W / aspect[0], "TXP_Max Bias", "dBm", "Bias(mA)")
    ax.plot(_num(m1, T), _num(m1, TX), color=C["purple"], lw=lw, label="[CH1_Max_TxP]")
    ax.plot(_num(m5, T), _num(m5, TX), color=C["navy"], lw=lw, label="[CH5_Max_TxP]")
    ax2.plot(_num(m1, T), _num(m1, B1), color=C["green"], lw=lw, label="[CH1_DDMI_Max_Bias]")
    lo = min(0, float(pd.concat([_num(m1, TX), _num(m5, TX)]).min() if len(m1) or len(m5) else 0))
    ax.set_ylim(int(lo) - (1 if lo < 0 else 0), 15)
    ax2.set_ylim(bottom=0)
    out["txp_max"] = _finish(fig, ax, ax2)

    for key, title, df, lab in (
            ("dw_o1", "Operational Current@dWL/dT Channel 1", o1, "[1_Opr_dW/dT]"),
            ("dw_o5", "Operational Current@dWL/dT Channel 5", o5, "[5_Opr_dW/dT]"),
            ("dw_m1", "Maximum Current@dWL/dT Channel 1", m1, "[1_Max_dW/dT]"),
            ("dw_m5", "Maximum Current@dWL/dT Channel 5", m5, "[5_Max_dW/dT]")):
        fig, ax, _ = _axes(W, W / aspect[1], title, "dWL/dT")
        ax.axhline(DWDT_MAX, color=C["spec"], lw=1.4, label="Max")
        ax.axhline(DWDT_MIN, color=C["spec"], lw=1.4, label="Min")
        ax.plot(_num(df, T), _num(df, dW), color=C["blue"], lw=lw, label=lab)
        ax.set_ylim(-0.6, 0.6)
        ax.set_yticks([-0.6, -0.4, -0.2, 0, 0.2, 0.4, 0.6])
        out[key] = _finish(fig, ax)
    return out


# --------------------------------------------------------------------------- 3. PPT
def _set_text(shape, text):
    """Replace a text box's text keeping the first run's font."""
    p = shape.text_frame.paragraphs[0]
    runs = p.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        p.add_run().text = text


def _set_lines(shape, lines, size_pt=None):
    """Put several lines into one paragraph separated by line breaks (like the template's Summary box)."""
    p = shape.text_frame.paragraphs[0]
    first = p.runs[0]._r if p.runs else None
    rpr = copy.deepcopy(first.find("{http://schemas.openxmlformats.org/drawingml/2006/main}rPr")) if first is not None else None
    for el in list(p._p):
        if el.tag.endswith("}r") or el.tag.endswith("}br") or el.tag.endswith("}fld"):
            p._p.remove(el)
    for i, line in enumerate(lines):
        if i:
            p.add_line_break()
        r = p.add_run()
        r.text = line
        if rpr is not None:
            r._r.insert(0, copy.deepcopy(rpr))
        if size_pt:
            r.font.size = Pt(size_pt)
    for extra in shape.text_frame.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)


def _clone_slide(prs, src):
    """Copy a template slide's non-picture shapes onto a new slide with the same layout."""
    s = prs.slides.add_slide(src.slide_layout)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    pics = []
    for sh in src.shapes:
        if sh.shape_type == 13:
            pics.append((sh.left, sh.top, sh.width, sh.height))
        else:
            s.shapes._spTree.append(copy.deepcopy(sh._element))
    # picture frames in TL, TR, BL, BR order
    pics.sort(key=lambda b: b[1])
    pics = sorted(pics[:2], key=lambda b: b[0]) + sorted(pics[2:], key=lambda b: b[0])
    return s, pics


def _shape(slide, name=None, startswith=None):
    for sh in slide.shapes:
        if sh.has_text_frame and ((name and sh.name == name) or (startswith and sh.text_frame.text.startswith(startswith))):
            return sh
    return None


def build_ppt(sns: list[SNData], conclusions: dict | None = None, report_date: str = "",
              title: str = "", purpose: str = "", condition: str = "", template=PPT_TEMPLATE,
              progress=None) -> bytes:
    prs = Presentation(template)
    slides = list(prs.slides)
    cover, summary, tpl_wl, tpl_dw = slides[0], slides[1], slides[2], slides[3]
    conclusions = conclusions or {}

    # cover
    for sh in cover.shapes:
        if not sh.has_text_frame:
            continue
        if re.match(r"^\s*\d{4}\s*/", sh.text_frame.text) and report_date:
            _set_text(sh, report_date)
        elif "Report" in sh.text_frame.text and title:
            _set_text(sh, title)

    # per-SN slides
    for i, d in enumerate(sns):
        charts = make_charts(d, aspect=(
            _ratio(tpl_wl), _ratio(tpl_dw)))
        wl_c, dw_c = conclusions.get(d.sn) or default_conclusions(d)
        for tpl, title_txt, keys, concl in (
                (tpl_wl, f"{d.sn}_Operational & Maximum Wavelength and TxP", ("wl1", "txp_opr", "wl5", "txp_max"), wl_c),
                (tpl_dw, f"{d.sn}_ Operational & Maximum dWL/dT", ("dw_o1", "dw_o5", "dw_m1", "dw_m5"), dw_c)):
            s, frames = _clone_slide(prs, tpl)
            _set_text(_title_box(s), title_txt)
            c = _shape(s, startswith="Conclusion")
            if c is not None:
                _set_text(c, f"Conclusion: {concl}")
            for k, (l, tp, w, h) in zip(keys, frames):
                s.shapes.add_picture(io.BytesIO(charts[k]), l, tp, w, h)
        if progress:
            progress((i + 1) / len(sns), d.sn)

    # summary slide
    lines = [f"Purpose: {purpose}".rstrip(), f"Test Condition: {condition}".rstrip(), "Conclusion:"]
    for d in sns:
        wl_c, dw_c = conclusions.get(d.sn) or default_conclusions(d)
        lines.append(f"   {d.sn}: {dw_c}" + ("" if wl_c == "Wavelength and TxP PASS" else f"; {wl_c}"))
    body = next((sh for sh in summary.shapes if sh.is_placeholder and sh.has_text_frame), None)
    if body is not None:
        _set_lines(body, lines, size_pt=24 if len(lines) <= 7 else 18 if len(lines) <= 11 else 14)

    # drop the two template SN slides
    sld = prs.slides._sldIdLst
    for tpl in (tpl_wl, tpl_dw):
        for sid in list(sld):
            if prs.part.related_part(sid.rId) is tpl.part:
                prs.part.drop_rel(sid.rId)
                sld.remove(sid)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _title_box(slide):
    """The purple-bar title text box (top-most text box)."""
    boxes = [sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    return min(boxes, key=lambda b: b.top)


def _ratio(slide):
    pics = [sh for sh in slide.shapes if sh.shape_type == 13]
    return sum(p.width / p.height for p in pics) / len(pics) if pics else 2.0


def ppt_filename(report_date: str = "") -> str:
    d = re.sub(r"\D", "", report_date) or date.today().strftime("%Y%m%d")
    return f"Sian3L_Mode_Hopping_Report_{d}.pptx"
